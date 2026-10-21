"""Stable relationships survive sorting; legacy row-based imports still work."""
import importlib.util
import pathlib
import unittest
import json
import tempfile
import openpyxl

spec = importlib.util.spec_from_file_location('framework_import', pathlib.Path(__file__).parents[1] / 'scripts/import-framework.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class FrameworkLinksTest(unittest.TestCase):
    def test_stable_ids_and_three_types_survive_reordering(self):
        rows = [(3, {'Evidence ID': 'UEH-0009', 'Supports': 'UEH-0004', 'Evaluates': 'UEH-0004'}),
                (9, {'Evidence ID': 'UEH-0004', 'Informs': 'UEH-0009; UEH-0009'})]
        self.assertEqual(module.framework_links(rows, {3: 'website-b', 9: 'website-a'}), [
            {'source': 'website-b', 'target': 'website-a', 'type': 'supports'},
            {'source': 'website-b', 'target': 'website-a', 'type': 'evaluates'},
            {'source': 'website-a', 'target': 'website-b', 'type': 'informs'}])

    def test_invalid_and_duplicate_ids_are_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Duplicate Evidence ID'):
            module.framework_links([(3, {'Evidence ID': 'A'}), (4, {'Evidence ID': 'A'})], {3: 'a', 4: 'b'})
        with self.assertRaisesRegex(ValueError, 'Invalid relationship'):
            module.framework_links([(3, {'Evidence ID': 'A', 'Informs': 'UEH-9999'})], {3: 'a'})

    def test_legacy_row_references_still_work(self):
        rows = [(3, {'Supporting links': '4', 'Direct links': '4'}), (4, {})]
        self.assertEqual(module.framework_links(rows, {3: 'a', 4: 'b'}), [
            {'source': 'a', 'target': 'b', 'type': 'supports'},
            {'source': 'a', 'target': 'b', 'type': 'links-to'}])

    def test_exact_file_metadata_and_intra_group_links_survive_reordering(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = pathlib.Path(directory)
            registry = directory / 'evidence.json'
            registry.write_text(json.dumps({'items': [{'id': 'group', 'title': 'Group', 'group': 'Group',
                'bucket': 'Research', 'tags': [], 'files': [
                    {'fileName': 'one.png', 'fileUrl': 'https://example.org/one.png'},
                    {'fileName': 'two.docx', 'fileUrl': 'https://example.org/two.docx'}]}]}))
            headers = ['Title', 'Item group name', 'Methods', 'Research', 'Analysis', 'Conclusion',
                       'Simulation', 'Strategy', 'Design', 'Evaluation'] + [''] * 45
            headers[10], headers[11], headers[12] = 'Vera', 'Image', 'Description'
            headers += ['Description/caption', 'Date', 'Source', 'Evidence ID', 'Informs', 'Supports', 'Evaluates']
            first, second = [None] * len(headers), [None] * len(headers)
            for row, title, identity, bucket, tag in [(first, 'one.png', 'UEH-A', 3, 11), (second, 'two.docx', 'UEH-B', 4, 12)]:
                row[0], row[1], row[10], row[bucket], row[tag], row[58] = title, 'Group', 1, 1, 1, identity
            first[59] = 'UEH-B'
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = 'Studio 2026'
            sheet.append(['Framework'])
            sheet.append(headers)
            sheet.append(first)
            sheet.append(second)
            workbook = directory / 'framework.xlsx'
            original = module.DATA
            try:
                module.DATA = registry
                for rows in [(first, second), (second, first)]:
                    for number, row in enumerate(rows, 3):
                        for column, value in enumerate(row, 1):
                            sheet.cell(number, column).value = value
                    book.save(workbook)
                    module.refresh(workbook)
                    payload = json.loads(registry.read_text())
                    files = payload['items'][0]['files']
                    self.assertEqual([file['framework']['id'] for file in files], ['UEH-A', 'UEH-B'])
                    self.assertIn('Image', files[0]['framework']['tags'])
                    self.assertNotIn('Image', files[1]['framework']['tags'])
                    self.assertEqual(payload['links'], [])
                    self.assertEqual(payload['fileLinks'], [{'source': 'UEH-A', 'target': 'UEH-B', 'type': 'informs'}])
                sheet.cell(4, 1).value = 'Updated title for the same file'
                book.save(workbook)
                module.refresh(workbook)
                files = json.loads(registry.read_text())['items'][0]['files']
                self.assertEqual(files[0]['framework']['id'], 'UEH-A')
                self.assertEqual(files[0]['fileUrl'], 'https://example.org/one.png')
                saved = registry.read_bytes()
                sheet.cell(4, 1).value = 'unmatched-file.png'
                sheet.cell(4, 59).value = 'UEH-new'
                book.save(workbook)
                with self.assertRaisesRegex(ValueError, 'exact attachment mapping'):
                    module.refresh(workbook)
                self.assertEqual(registry.read_bytes(), saved, 'Failed imports must leave the public snapshot intact')
            finally:
                module.DATA = original


if __name__ == '__main__':
    unittest.main()
